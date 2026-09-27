"""CSV formula-injection escaping and HTML (XSS) escaping in generated reports."""
from __future__ import annotations

import csv
import io
import json
import subprocess
import sys
from pathlib import Path

import pytest

from jiangkit.security.csvsafe import SafeDictWriter, SafeWriter, escape_cell, escape_dataframe

from ._util import jiang

XSS = '<script>alert("x")</script><img src=x onerror=alert(1)>'


@pytest.mark.parametrize("cell", ["=1+1", "=HYPERLINK(\"http://evil\")", "+cmd|' /C calc'!A0", "-2+3=1",
                                  "@SUM(A1:A9)", "\t=1", "\r=1", "＝1+1"])
def test_formulas_escaped(cell):
    assert escape_cell(cell).startswith("'")


@pytest.mark.parametrize("cell", ["@handle", "-3.5", "+886912345678", "- bullet point", "hello", "", 42, None])
def test_harmless_values_untouched(cell):
    assert escape_cell(cell) == cell


def test_safe_writers():
    buf = io.StringIO()
    SafeWriter(buf).writerow(["=cmd", "ok"])
    w = SafeDictWriter(buf, fieldnames=["a"])
    w.writerow({"a": "@SUM(1)"})
    rows = list(csv.reader(io.StringIO(buf.getvalue())))
    assert rows[0][0] == "'=cmd" and rows[1][0] == "'@SUM(1)"


def test_escape_dataframe():
    import pandas as pd
    df = escape_dataframe(pd.DataFrame({"=h": ["=1", "x"], "n": [1, 2]}))
    assert list(df.columns)[0] == "'=h" and df.iloc[0, 0] == "'=1" and df.iloc[0, 1] == 1


def test_data_tools_csv_export_escapes(tmp_path):
    src = tmp_path / "in.json"
    src.write_text(json.dumps([{"name": "=HYPERLINK(\"http://evil\",\"x\")", "handle": "@mei"}]), "utf-8")
    out = tmp_path / "out.csv"
    r = jiang("data", "data", "convert", src, "-o", out)
    assert r.returncode == 0, r.stderr
    rows = list(csv.DictReader(out.open(encoding="utf-8-sig")))
    assert rows[0]["name"].startswith("'=") and rows[0]["handle"] == "@mei"


def test_data_tools_xlsx_export_escapes(tmp_path):
    import openpyxl
    src = tmp_path / "in.csv"
    src.write_text("a\n=1+2\n", "utf-8")
    out = tmp_path / "out.xlsx"
    assert jiang("data", "data", "convert", src, "-o", out).returncode == 0
    ws = openpyxl.load_workbook(out).active
    assert ws["A2"].value == "'=1+2" and ws["A2"].data_type == "s"


def test_dream_diff_html_escapes(tmp_path):
    from jiangkit.experimental import dream_diff as dd
    old, new = f"The price is 10 dollars. {XSS}", f"The price is 12 dollars. {XSS} We must pay."
    ch = dd.semantic_diff(old, new)
    html = dd.html_report(ch, dd.summarize(ch, old, new), XSS, "b<i>.txt")
    assert "<script>" not in html and "onerror=alert" not in html.replace("onerror=alert(1)&gt;", "")
    assert "&lt;script&gt;" in html and "Content-Security-Policy" in html


def test_code_weather_html_escapes(tmp_path):
    repo = tmp_path / "repo"
    (repo / "pkg").mkdir(parents=True)
    (repo / "pkg" / "x<img src=y onerror=alert(1)>.py").write_text("def f():\n    return 1  # TODO <script>\n")
    out = tmp_path / "w.html"
    r = jiang("experimental", "code_weather", repo, "--html", out)
    assert r.returncode == 0, r.stderr
    html = out.read_text()
    assert "<img src=y" not in html and "&lt;img src=y" in html and "Content-Security-Policy" in html


def test_no_unsafe_yaml_or_pickle_in_sources():
    src = Path(__file__).resolve().parents[2] / "src" / "jiangkit"
    offenders = []
    for f in src.rglob("*.py"):
        t = f.read_text(encoding="utf-8")
        if "yaml.load(" in t and "SafeLoader" not in t:
            offenders.append(f"{f}: yaml.load without SafeLoader")
        if "yaml.unsafe_load" in t or "pickle.load" in t or "import pickle" in t or "marshal.loads" in t:
            offenders.append(f"{f}: pickle/marshal")
        if "torch.load(" in t and "weights_only=True" not in t:
            offenders.append(f"{f}: torch.load without weights_only")
    assert not offenders, offenders


def test_openpyxl_uses_defusedxml():
    import openpyxl.xml
    assert openpyxl.xml.DEFUSEDXML, "install defusedxml (in the [data] extra) to parse untrusted XLSX safely"
